#!/usr/bin/env python3
"""
build_slide_1.py
Constructs Slide 1 of the Final Defense Presentation based on write_up/Template .pptx.
Applies visual hierarchy, professional typography (Gill Sans MT / Deep Navy #002060),
and zero-bleeding safe margins.
"""

import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

TEMPLATE_PATH = "write_up/Template .pptx"
OUTPUT_PATH = "write_up/Project Final Defense Slides_FINAL.pptx"

def set_para(p, text, font_name="Gill Sans MT", size_pt=11, bold=False, italic=False, color_rgb=None, space_after_pt=2, line_spacing=1.1, align=None):
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(size_pt)
    p.font.bold = bold
    p.font.italic = italic
    p.space_after = Pt(space_after_pt)
    p.line_spacing = line_spacing
    if color_rgb:
        p.font.color.rgb = color_rgb
    if align:
        p.alignment = align

def main():
    # Start fresh from the original template
    shutil.copy(TEMPLATE_PATH, OUTPUT_PATH)
    prs = pptx.Presentation(OUTPUT_PATH)
    s1 = prs.slides[0]

    # Colors
    c_navy   = RGBColor(0, 32, 96)      # #002060
    c_dark   = RGBColor(15, 23, 42)     # #0F172A
    c_slate  = RGBColor(71, 85, 105)    # #475569
    c_muted  = RGBColor(100, 116, 139)  # #64748B

    # Find Rectangle 3 (Right content panel)
    target_shape = None
    for shape in s1.shapes:
        if shape.name == "Rectangle 3":
            target_shape = shape
            break
        elif shape.name == "Title 1":
            # Clear empty title placeholder
            if shape.has_text_frame:
                shape.text_frame.clear()

    if target_shape and target_shape.has_text_frame:
        # Position Rectangle 3 comfortably inside safe boundaries
        target_shape.left = Inches(6.55)
        target_shape.top = Inches(0.50)
        target_shape.width = Inches(6.30)
        target_shape.height = Inches(6.50)
        
        tf = target_shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.15)
        tf.margin_bottom = Inches(0.15)
        tf.clear()

        # 1. Institutional Header
        p1 = tf.paragraphs[0]
        set_para(p1, "FACULTY OF ENGINEERING", size_pt=12, bold=True, color_rgb=c_navy, space_after_pt=1)
        
        p2 = tf.add_paragraph()
        set_para(p2, "DEPARTMENT OF COMPUTER ENGINEERING", size_pt=12, bold=True, color_rgb=c_navy, space_after_pt=12)

        # 2. Main Title (High Visual Hierarchy)
        p3 = tf.add_paragraph()
        set_para(p3, "DEVELOPMENT AND IMPLEMENTATION OF A ROBOTIC APPLICATION FOR HUMAN INTENTION RECOGNITION USING MOTION AND HAND GESTURE", 
                 size_pt=14.5, bold=True, color_rgb=c_dark, space_after_pt=10, line_spacing=1.15)

        # 3. Subtitle / Academic Degree
        p4 = tf.add_paragraph()
        set_para(p4, "A Final Project Report Submitted in Partial Fulfillment of the Requirements for the Degree of Bachelor of Science in Computer Engineering",
                 size_pt=10.5, italic=True, color_rgb=c_slate, space_after_pt=16, line_spacing=1.15)

        # 4. Presenters Block
        p5 = tf.add_paragraph()
        set_para(p5, "PRESENTED BY:", size_pt=10.5, bold=True, color_rgb=c_navy, space_after_pt=2)
        
        p6 = tf.add_paragraph()
        set_para(p6, "ELEANA OSEI OWUSU  (Index: 4121230024)", size_pt=11, bold=True, color_rgb=c_dark, space_after_pt=1)
        
        p7 = tf.add_paragraph()
        set_para(p7, "JOEL NII ADJETEY AHULU  (Index: 4121230020)", size_pt=11, bold=True, color_rgb=c_dark, space_after_pt=14)

        # 5. Supervision & Date Block
        p8 = tf.add_paragraph()
        set_para(p8, "SUPERVISOR: MR. MICHEAL XENYA", size_pt=11, bold=True, color_rgb=c_navy, space_after_pt=2)
        
        p9 = tf.add_paragraph()
        set_para(p9, "ACADEMIC YEAR: 2025/2026", size_pt=10, color_rgb=c_muted, space_after_pt=0)

    prs.save(OUTPUT_PATH)
    print(f"Slide 1 successfully built and saved to: {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
