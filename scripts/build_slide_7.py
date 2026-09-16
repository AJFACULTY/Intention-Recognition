#!/usr/bin/env python3
"""
build_slide_7.py — Formal Academic Slide 7 (6-Phase Engineering Methodology)
- 6-Phase Engineering Research & Implementation Methodology.
- Open Editorial Typography: cardless, breathable, high-contrast.
- 2 rows x 3 columns symmetric layout.
- Sleek sequence flow arrows aligned with Phase badges.
- Single-line title with generous vertical breathing room.
- Generous bottom clearance (>0.75" above footer bar).
- Dignified GCTU 2-tone palette: Navy (#002060), Gold (#B8860B), Charcoal (#1E293B).
"""

import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = "write_up/Project Final Defense Slides_FINAL.pptx"

def set_para(p, text, font_name="Calisto MT", size_pt=14, bold=False, italic=False, color_rgb=None, space_before_pt=0, space_after_pt=2, line_spacing=1.14, align=PP_ALIGN.LEFT):
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(size_pt)
    p.font.bold = bold
    p.font.italic = italic
    p.space_before = Pt(space_before_pt)
    p.space_after = Pt(space_after_pt)
    p.line_spacing = line_spacing
    if color_rgb:
        p.font.color.rgb = color_rgb
    p.alignment = align

def add_phase_column(slide, name, left, top, width, height, phase_num, title, items, c_navy, c_gold, c_dark):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tx.name = name
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.02)

    # 1. Phase Badge (e.g. PHASE 01) in Gold
    p_badge = tf.paragraphs[0]
    set_para(p_badge, phase_num, font_name="Calisto MT", size_pt=10, bold=True, color_rgb=c_gold, space_after_pt=1, line_spacing=1.10)

    # 2. Phase Title in Bold Navy
    p_title = tf.add_paragraph()
    set_para(p_title, title, font_name="Calisto MT", size_pt=12, bold=True, color_rgb=c_navy, space_after_pt=4, line_spacing=1.12)

    # 3. Bullets in Charcoal
    for i, item in enumerate(items):
        p_item = tf.add_paragraph()
        space_after = 3 if i < len(items) - 1 else 0
        set_para(p_item, f"• {item}", font_name="Calisto MT", size_pt=10, bold=False, color_rgb=c_dark, space_after_pt=space_after, line_spacing=1.14)

def main():
    prs = pptx.Presentation(PPTX_PATH)

    if len(prs.slides) < 7:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
    else:
        slide = prs.slides[6]

    # Theme colors
    c_navy   = RGBColor(0, 32, 96)       # #002060 GCTU Deep Navy
    c_dark   = RGBColor(30, 41, 59)      # #1E293B High-Contrast Charcoal
    c_gold   = RGBColor(184, 134, 11)    # #B8860B Accent Gold
    c_arrow  = RGBColor(184, 134, 11)    # Gold connector arrows
    c_slate  = RGBColor(100, 116, 139)   # Slate gray

    # 1. Slide Title (Single Line Fit at 24 pt Bold Navy)
    for shape in slide.shapes:
        if shape.name == "Title 1" and shape.has_text_frame:
            shape.left = Inches(0.66)
            shape.top = Inches(0.35)
            shape.width = Inches(12.00)
            shape.height = Inches(0.60)
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            set_para(p, "05. 6-PHASE RESEARCH & IMPLEMENTATION METHODOLOGY", font_name="Calisto MT", size_pt=24, bold=True, color_rgb=c_navy, align=PP_ALIGN.CENTER)

        if shape.name == "Content Placeholder 2" and shape.has_text_frame:
            shape.text_frame.clear()

    # Clean existing custom shapes on Slide 7
    shapes_to_remove = [s for s in slide.shapes if s.name.startswith("Meth_")]
    for s in shapes_to_remove:
        el = s._element
        el.getparent().remove(el)

    # 2. Section Subhead: Left-aligned Gold Eyebrow
    tx_sub = slide.shapes.add_textbox(Inches(0.85), Inches(1.05), Inches(11.63), Inches(0.25))
    tx_sub.name = "Meth_Subhead"
    tf_sub = tx_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    set_para(p_sub, "STRUCTURED RESEARCH & IMPLEMENTATION LIFECYCLE FROM DESIGN TO VALIDATION", font_name="Calisto MT", size_pt=11, bold=True, color_rgb=c_gold, align=PP_ALIGN.LEFT)

    # Coordinates
    # 3 Columns: Lefts: 0.85", 4.90", 8.95" | Width: 3.53"
    col_lefts = [Inches(0.85), Inches(4.90), Inches(8.95)]
    col_w = Inches(3.53)
    row0_top = Inches(1.38)
    row_h = Inches(2.05)
    row1_top = Inches(3.85)

    # Sequence Arrow positions aligned with Phase Badges (Phase 1 -> 2, Phase 2 -> 3)
    # Gap 0-1: 4.38" to 4.90" -> Arrow Left: 4.48", Width: 0.32"
    # Gap 1-2: 8.43" to 8.95" -> Arrow Left: 8.53", Width: 0.32"
    arrow_lefts = [Inches(4.48), Inches(8.53)]
    arrow_w = Inches(0.32)
    arrow_h = Inches(0.13)

    # 3. Row 0: Phases 1 to 3
    phases_row0 = [
        {
            "num": "PHASE 01",
            "title": "Mechatronics & Embedded Hardware",
            "items": [
                "Built Yahboom 4WD mobile chassis with DC encoder motors and 2-DOF camera gimbal.",
                "Integrated Raspberry Pi 5 (8GB) edge compute node and ESP32-S3 micro-ROS co-processor.",
                "Implemented decoupled dual-bus power regulation (7.4V Li-ion) isolating logic from motor spikes."
            ]
        },
        {
            "num": "PHASE 02",
            "title": "Data Acquisition & Calibration",
            "items": [
                "Captured and balanced custom 6,000-sample dataset directly through onboard camera.",
                "Mapped 6 operational gesture classes: STOP, GO, FOLLOW, LEFT, RIGHT, BACK.",
                "Calibrated across varying distances (1.0m to 2.5m) under natural indoor lighting."
            ]
        },
        {
            "num": "PHASE 03",
            "title": "Feature Engineering & Edge AI",
            "items": [
                "Extracted 21 MediaPipe hand landmarks and computed 19-D geometric invariant features.",
                "Trained lightweight PyTorch MLP neural classifier (46 KB, sub-2ms CPU inference).",
                "Implemented 10-frame LSTM recurrent network for short-horizon human motion prediction."
            ]
        }
    ]

    for idx, p in enumerate(phases_row0):
        add_phase_column(slide, f"Meth_P{idx+1}", col_lefts[idx], row0_top, col_w, row_h, p["num"], p["title"], p["items"], c_navy, c_gold, c_dark)

    # Sequence Arrows for Row 0 (Aligned vertically with the PHASE badges at Y = 1.42")
    for idx, a_left in enumerate(arrow_lefts):
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, a_left, Inches(1.42), arrow_w, arrow_h)
        arr.name = f"Meth_Arrow_R0_{idx}"
        arr.fill.solid()
        arr.fill.fore_color.rgb = c_arrow
        arr.line.color.rgb = c_arrow

    # 4. Transitional Flow Ribbon between Row 0 and Row 1
    tx_trans = slide.shapes.add_textbox(Inches(1.50), Inches(3.50), Inches(10.33), Inches(0.26))
    tx_trans.name = "Meth_Transition"
    tf_trans = tx_trans.text_frame
    tf_trans.word_wrap = True
    tf_trans.margin_left = tf_trans.margin_right = tf_trans.margin_top = tf_trans.margin_bottom = 0
    p_trans = tf_trans.paragraphs[0]
    set_para(p_trans, "──  SEQUENTIAL PROGRESSION TO SUPERVISORY MIDDLEWARE & NAVIGATION INTEGRATION  ──", font_name="Calisto MT", size_pt=9.5, bold=True, color_rgb=c_slate, align=PP_ALIGN.CENTER)

    # 5. Row 1: Phases 4 to 6
    phases_row1 = [
        {
            "num": "PHASE 04",
            "title": "Supervisory Control & Safety",
            "items": [
                "Established centralized ROS 2 Humble node graph on dedicated communication Domain 20.",
                "Engineered brain_node state machine with 5-frame rolling consensus and visual servoing.",
                "Enforced priority velocity multiplexing (twist_mux) with 0.36m reactive laser safety stop."
            ]
        },
        {
            "num": "PHASE 05",
            "title": "Metric SLAM & Autonomous Navigation",
            "items": [
                "Generated high-fidelity 2D occupancy grid maps (5cm resolution) via slam_toolbox.",
                "Fused wheel odometry and 6-axis IMU through Extended Kalman Filter (EKF) localization.",
                "Configured Nav2 global and local costmaps with 0.25m obstacle inflation padding."
            ]
        },
        {
            "num": "PHASE 06",
            "title": "Empirical Benchmarking & Validation",
            "items": [
                "Quantified end-to-end component latency budget at 74.2 ms (well under 150 ms threshold).",
                "Validated closed-loop control across 180 physical trials (96.67% real-world accuracy).",
                "Verified collaborative safety compliance under ISO 15066 (<0.36m deterministic halt)."
            ]
        }
    ]

    for idx, p in enumerate(phases_row1):
        add_phase_column(slide, f"Meth_P{idx+4}", col_lefts[idx], row1_top, col_w, row_h, p["num"], p["title"], p["items"], c_navy, c_gold, c_dark)

    # Sequence Arrows for Row 1 (Aligned vertically with the PHASE badges at Y = 3.89")
    for idx, a_left in enumerate(arrow_lefts):
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, a_left, Inches(3.89), arrow_w, arrow_h)
        arr.name = f"Meth_Arrow_R1_{idx}"
        arr.fill.solid()
        arr.fill.fore_color.rgb = c_arrow
        arr.line.color.rgb = c_arrow

    prs.save(PPTX_PATH)
    print(f"Slide 7 (6-Phase Methodology) successfully updated in {PPTX_PATH}")

if __name__ == "__main__":
    main()
