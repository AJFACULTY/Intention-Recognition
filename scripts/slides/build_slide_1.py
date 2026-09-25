#!/usr/bin/env python3
"""
build_slide_1.py — Fixed Professional Slide 1
Maintains Rectangle 3 at full height (0 to 7.5 inches) covering the right panel cleanly.
Sets text to brilliant pure white (#FFFFFF) and gold (#FDB813) on the deep navy (#002060) background.
Centers text with high visual hierarchy, clean line spacing, and correct student IDs.
"""

import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE_PATH = "write_up/Template .pptx"
OUTPUT_PATH = "write_up/Project Final Defense Slides_FINAL.pptx"

def set_para(p, text, font_name="Calisto MT", size_pt=14, bold=False, italic=False, color_rgb=None, space_after_pt=4, line_spacing=1.15, align=PP_ALIGN.CENTER):
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(size_pt)
    p.font.bold = bold
    p.font.italic = italic
    p.space_after = Pt(space_after_pt)
    p.line_spacing = line_spacing
    if color_rgb:
        p.font.color.rgb = color_rgb
    p.alignment = align

def main():
    shutil.copy(TEMPLATE_PATH, OUTPUT_PATH)
    prs = pptx.Presentation(OUTPUT_PATH)
    s1 = prs.slides[0]

    # Colors on Dark Navy (#002060)
    c_white  = RGBColor(255, 255, 255)  # #FFFFFF
    c_gold   = RGBColor(253, 184, 19)   # #FDB813 GCTU Gold
    c_sub    = RGBColor(226, 232, 240)  # #E2E8F0 Light Slate

    # Clean Title 1 placeholder if present
    for shape in s1.shapes:
        if shape.name == "Title 1" and shape.has_text_frame:
            shape.text_frame.clear()

    # Find Rectangle 3 (Right full panel)
    for shape in s1.shapes:
        if shape.name == "Rectangle 3" and shape.has_text_frame:
            # Keep original full panel bounds: covers from left=6.48 to right edge, top=0 to bottom=7.5
            shape.left = Inches(6.476)
            shape.top = Inches(0.0)
            shape.width = Inches(6.857)
            shape.height = Inches(7.50)
            
            # Ensure solid navy fill
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0, 32, 96) # #002060
            
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.40)
            tf.margin_right = Inches(0.40)
            tf.margin_top = Inches(0.35)
            tf.margin_bottom = Inches(0.35)
            tf.clear()

            # 1. Institutional Header
            p1 = tf.paragraphs[0]
            set_para(p1, "FACULTY OF ENGINEERING", size_pt=14, bold=True, color_rgb=c_gold, space_after_pt=1)
            
            p2 = tf.add_paragraph()
            set_para(p2, "DEPARTMENT OF COMPUTER ENGINEERING", size_pt=13, bold=True, color_rgb=c_sub, space_after_pt=14)

            # 2. Main Title (Dominant Visual Focus)
            p3 = tf.add_paragraph()
            set_para(p3, "DEVELOPMENT AND IMPLEMENTATION OF A ROBOTIC APPLICATION FOR HUMAN INTENTION RECOGNITION USING MOTION AND HAND GESTURE", 
                     size_pt=18, bold=True, color_rgb=c_white, space_after_pt=12, line_spacing=1.15)

            # 3. Subtitle / Academic Degree
            p4 = tf.add_paragraph()
            set_para(p4, "A Final Project Report Submitted in Partial Fulfillment of the Requirements for the Degree of Bachelor of Science in Computer Engineering",
                     size_pt=12, italic=True, color_rgb=c_sub, space_after_pt=18, line_spacing=1.15)

            # 4. Presenters Block
            p5 = tf.add_paragraph()
            set_para(p5, "PRESENTATION BY:", size_pt=12, bold=True, color_rgb=c_gold, space_after_pt=2)
            
            p6 = tf.add_paragraph()
            set_para(p6, "ELEANA OSEI OWUSU — 4121230024", size_pt=13.5, bold=True, color_rgb=c_white, space_after_pt=2)
            
            p7 = tf.add_paragraph()
            set_para(p7, "JOEL NII ADJETEY AHULU — 4121230020", size_pt=13.5, bold=True, color_rgb=c_white, space_after_pt=16)

            # 5. Supervision & Academic Year
            p8 = tf.add_paragraph()
            set_para(p8, "SUPERVISOR: MR. MICHEAL XENYA", size_pt=13, bold=True, color_rgb=c_gold, space_after_pt=2)
            
            p9 = tf.add_paragraph()
            set_para(p9, "ACADEMIC YEAR: 2025/2026", size_pt=11.5, color_rgb=c_sub, space_after_pt=0)

    prs.save(OUTPUT_PATH)
    print(f"Slide 1 rebuilt: {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
