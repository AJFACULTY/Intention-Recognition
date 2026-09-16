#!/usr/bin/env python3
"""
build_slide_5.py — Option 1: Open Editorial Typography (Cardless Layout)
- Slide Title: 03. PROJECT OBJECTIVES
- General Objective: Elegant open callout with Gold eyebrow header and subtle divider rule.
- Specific Objectives: Sleek 2-column x 3-row open editorial layout.
- Zero boxes, zero border cards — clean, modern, and breathable.
"""

import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR

PPTX_PATH = "write_up/Project Final Defense Slides_FINAL.pptx"

def set_para(p, text, font_name="Calisto MT", size_pt=14, bold=False, italic=False, color_rgb=None, space_before_pt=0, space_after_pt=2, line_spacing=1.15, align=PP_ALIGN.LEFT):
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(size_pt)
    p.font.bold = bold
    p.font.italic = italic
    p.space_before = Pt(space_before_pt)
    p.space_after = Pt(space_after_pt)
    p.line_spacing = line_spacing
    if color_rgb:
        p.font.color.rgb = color_rgb
    p.alignment = align

def main():
    prs = pptx.Presentation(PPTX_PATH)
    slide = prs.slides[4]

    # Theme colors matching Slides 1-4
    c_navy   = RGBColor(0, 32, 96)       # #002060 GCTU Deep Navy
    c_body   = RGBColor(51, 65, 85)      # #334155 Slate Charcoal
    c_gold   = RGBColor(184, 134, 11)    # #B8860B Accent Gold
    c_border = RGBColor(203, 213, 225)   # #CBD5E1 Subtle Slate

    # 1. Slide Title (28 pt Bold Navy)
    for shape in slide.shapes:
        if shape.name == "Title 1" and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            set_para(p, "03. PROJECT OBJECTIVES", font_name="Calisto MT", size_pt=28, bold=True, color_rgb=c_navy, align=PP_ALIGN.CENTER)

        if shape.name == "Content Placeholder 2" and shape.has_text_frame:
            shape.text_frame.clear()

    # Remove all previous custom shapes on Slide 5
    shapes_to_remove = [s for s in slide.shapes if s.name.startswith("Obj_")]
    for s in shapes_to_remove:
        el = s._element
        el.getparent().remove(el)

    # 2. General Objective: Open Editorial Callout (No Box)
    # Left=1.00", Top=1.25", Width=11.33", Height=1.05"
    tx_gen = slide.shapes.add_textbox(Inches(1.00), Inches(1.25), Inches(11.33), Inches(1.05))
    tx_gen.name = "Obj_Gen_Text"
    tf_gen = tx_gen.text_frame
    tf_gen.word_wrap = True
    tf_gen.margin_left = Inches(0.05)
    tf_gen.margin_right = Inches(0.05)
    tf_gen.margin_top = Inches(0.02)
    tf_gen.margin_bottom = Inches(0.02)

    p_eyebrow = tf_gen.paragraphs[0]
    set_para(p_eyebrow, "GENERAL OBJECTIVE", font_name="Calisto MT", size_pt=13, bold=True, color_rgb=c_gold, space_after_pt=4, align=PP_ALIGN.CENTER)

    p_gen = tf_gen.add_paragraph()
    set_para(p_gen, 
             "To design, implement, and evaluate an edge-deployed, real-time human intention recognition and autonomous navigation system powered by ROS 2 on an embedded mobile robot.",
             font_name="Calisto MT", size_pt=15, bold=False, color_rgb=RGBColor(30, 41, 59), line_spacing=1.18, align=PP_ALIGN.CENTER)

    # Subtle horizontal dividing rule between General & Specific Objectives
    # Y = 2.45", from X=2.50" to X=10.83"
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.50), Inches(2.45), Inches(10.83), Inches(2.45))
    line.name = "Obj_Divider_Line"
    line.line.color.rgb = c_border
    line.line.width = Pt(0.8)

    # 3. Specific Objectives: 2 Columns x 3 Rows Open Editorial Layout
    # Col 1: Left=1.00", Width=5.35"
    # Col 2: Left=6.98", Width=5.35"
    # Rows: 2.65", 3.82", 4.99"
    lefts = [Inches(1.00), Inches(6.98)]
    row_tops = [Inches(2.65), Inches(3.82), Inches(4.99)]
    item_w = Inches(5.35)
    item_h = Inches(1.05)

    specific_objectives = [
        {
            "num": "01",
            "title": "Invariant Vision Pipeline",
            "desc": "Compute 19 scale-invariant geometric features from 21 hand landmarks for real-time onboard MLP classification."
        },
        {
            "num": "02",
            "title": "Balanced 6,000-Sample Dataset",
            "desc": "Capture and balance a custom 6-class dataset directly through the onboard camera to eliminate perspective domain shift."
        },
        {
            "num": "03",
            "title": "LSTM Trajectory Prediction",
            "desc": "Forecast short-horizon operator motion trajectories from continuous body pose landmarks using recurrent neural networks."
        },
        {
            "num": "04",
            "title": "Supervisory Brain Node",
            "desc": "Execute a deterministic finite state machine with 5-frame rolling consensus, active operator lock, and visual servoing."
        },
        {
            "num": "05",
            "title": "2D LiDAR SLAM & Nav2",
            "desc": "Deploy real-time Cartographer mapping and the Nav2 navigation stack for socially aware, collision-free path execution."
        },
        {
            "num": "06",
            "title": "Empirical Safety Benchmarking",
            "desc": "Quantify end-to-end latency (132 ms), physical classification accuracy (96.7%), and compliance with ISO 15066 safety margins."
        }
    ]

    # Map items to 2 columns:
    # Column 0: items 0, 1, 2
    # Column 1: items 3, 4, 5
    for idx, item in enumerate(specific_objectives):
        col = 0 if idx < 3 else 1
        row = idx % 3

        box = slide.shapes.add_textbox(lefts[col], row_tops[row], item_w, item_h)
        box.name = f"Obj_Item_{idx}"
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.05)
        tf.margin_right = Inches(0.05)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)

        # Title Paragraph with Numeral
        p_title = tf.paragraphs[0]
        p_title.text = f"{item['num']}.  {item['title']}"
        p_title.font.name = "Calisto MT"
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = c_navy
        p_title.space_after = Pt(2)
        p_title.alignment = PP_ALIGN.LEFT

        # Description Paragraph
        p_desc = tf.add_paragraph()
        set_para(p_desc, item["desc"], font_name="Calisto MT", size_pt=12, color_rgb=c_body, line_spacing=1.15, align=PP_ALIGN.LEFT)

    prs.save(PPTX_PATH)
    print(f"Slide 5 rebuilt with Open Editorial Layout in {PPTX_PATH}")

if __name__ == "__main__":
    main()
